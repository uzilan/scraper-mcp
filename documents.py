import json
from pathlib import Path

import chromadb

import indexing
import namespaces
import openapi
from parsing import chunk


def extract_text_pdf(path: Path) -> str:
    from pypdf import PdfReader
    reader = PdfReader(path)
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def extract_text_docx(path: Path) -> str:
    from docx import Document
    doc = Document(path)
    return "\n".join(p.text for p in doc.paragraphs if p.text)


def extract_text_xlsx(path: Path) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    lines = []
    for sheet in wb.worksheets:
        lines.append(f"# {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                lines.append("\t".join(cells))
    return "\n".join(lines)


def index_file(collection: chromadb.Collection, namespace: str, path: Path) -> str:
    import yaml as _yaml
    source_url = f"file://{namespace}/{path.name}"
    suffix = path.suffix.lower()

    if suffix == ".doc":
        return f".doc is not supported — convert to .docx first."

    if suffix in (".json", ".yaml", ".yml"):
        text = path.read_text(encoding="utf-8")
        data = json.loads(text) if suffix == ".json" else _yaml.safe_load(text)
        if isinstance(data, dict) and "paths" in data and ("openapi" in data or "swagger" in data):
            docs, ids, metas = openapi.openapi_documents(data, source_url)
            ok = indexing.upsert(collection, docs, ids, metas, source_url)
            return f"Indexed {len(docs)} operations from {path.name}" if ok else f"No content extracted from {path.name}"
        text_content = json.dumps(data, indent=2) if suffix == ".json" else text
    elif suffix == ".pdf":
        text_content = extract_text_pdf(path)
    elif suffix == ".docx":
        text_content = extract_text_docx(path)
    elif suffix == ".xlsx":
        text_content = extract_text_xlsx(path)
    elif suffix in (".txt", ".md"):
        text_content = path.read_text(encoding="utf-8")
    else:
        return f"Unsupported file type: {suffix}"

    chunks_list = chunk(text_content)
    ids = [f"{source_url}::{i}" for i in range(len(chunks_list))]
    metas = [{"source_url": source_url} for _ in chunks_list]
    ok = indexing.upsert(collection, chunks_list, ids, metas, source_url)
    return f"Indexed {len(chunks_list)} chunks from {path.name}" if ok else f"No content extracted from {path.name}"


def list_documents(namespace: str) -> list[dict]:
    import mimetypes
    folder = namespaces.UPLOADS_PATH / namespace
    if not folder.exists():
        return []
    return [
        {
            "name": f.name,
            "size": f.stat().st_size,
            "content_type": mimetypes.guess_type(f.name)[0] or "application/octet-stream",
        }
        for f in sorted(folder.iterdir())
        if f.is_file()
    ]


def delete_document(collection: chromadb.Collection, namespace: str, filename: str) -> str:
    path = namespaces.UPLOADS_PATH / namespace / filename
    source_url = f"file://{namespace}/{filename}"
    existing = collection.get(where={"source_url": source_url})
    if existing["ids"]:
        collection.delete(ids=existing["ids"])
    if path.exists():
        path.unlink()
        return f"Deleted {filename}"
    return f"File {filename} not found"
